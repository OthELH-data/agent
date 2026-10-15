#!/usr/bin/env python3
"""One-time builder for the Vito/Sparta/Whale IPT registry.

Runtime code never parses these workbooks or presentations.  Re-run this tool only when the source
charts change, review entries marked ``needs_review``, then commit ``configs/ipt_registry.yaml``.
"""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
SOURCE_ROOT = Path(
    "/Users/harshakolachina/Documents/Projects/Shell JV Audit/Data Files/"
    "IPT Org Charts - Vito, Sparta & Whale"
)
OUTPUT = ROOT / "configs" / "ipt_registry.yaml"

_HEADING_WORDS = re.compile(
    r"\b(ipt|team|operations|readiness|wells|subsurface|topsides|surf|construction|finance|imit|"
    r"or&a|manager|lead|engineering|project|support|drilling|completion|commissioning|hsse|"
    r"commercial|contract|procurement|civil|marine|hull|facility|maintenance)\b",
    re.I,
)
_NON_NAME = re.compile(r"\b(team|chart|manager|management|engineering|operations|project|support|"
                       r"construction|finance|wells|subsurface|topsides|contract|commissioning|"
                       r"readiness|lead|specialist|engineer|supervisor|advisor|planner|technician)\b", re.I)


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _tokens(value: str) -> list[str]:
    return re.findall(r"[A-Za-z][A-Za-z'’-]*", _clean(value))


def _looks_like_name(value: str) -> bool:
    text = _clean(value)
    if not text or any(char.isdigit() for char in text) or _NON_NAME.search(text):
        return False
    tokens = _tokens(text)
    if not 2 <= len(tokens) <= 7:
        return False
    return "," in text or sum(token[:1].isupper() or token.isupper() for token in tokens) >= 2


def _asset(path: Path) -> str:
    lower = path.name.casefold()
    return next(asset for asset in ("Vito", "Sparta", "Whale") if asset.casefold() in lower)


def _year(path: Path) -> str:
    return next(part for part in path.parts if re.fullmatch(r"20\d{2}", part))


def _source_record(path: Path) -> dict[str, Any]:
    dates = re.findall(r"(?:19|20)\d{2}[-_/]\d{1,2}[-_/]\d{1,2}|(?:19|20)\d{2}", path.name)
    return {"file": path.name, "effective_date": dates[-1] if dates else ""}


def _add(store: dict, *, year: str, asset: str, name: str, title: str = "",
         context: str = "", employee_id: str = "", source: str, source_ref: str,
         review_status: str = "confirmed") -> None:
    name = _clean(name)
    if not _looks_like_name(name):
        return
    key = re.sub(r"[^a-z0-9]", "", name.casefold())
    bucket = store.setdefault(year, {}).setdefault(asset, {"sources": [], "people": {}})
    person = bucket["people"].setdefault(key, {
        "name": name, "aliases": [], "employee_id": employee_id or None,
        "registry_title": title or None, "chart_context": [], "ipt": True,
        "source_refs": [], "review_status": review_status,
    })
    if name != person["name"] and name not in person["aliases"]:
        person["aliases"].append(name)
    tokens = name.split()
    if len(tokens) >= 3 and (
        (len(tokens[0]) == 1 and tokens[0].isalpha())
        or (tokens[0].isupper() and 2 <= len(tokens[0]) <= 6)
    ):
        stripped = " ".join(tokens[1:])
        if _looks_like_name(stripped) and stripped not in person["aliases"]:
            person["aliases"].append(stripped)
    if title and not person.get("registry_title"):
        person["registry_title"] = title
    if employee_id and not person.get("employee_id"):
        person["employee_id"] = employee_id
    if context and context not in person["chart_context"]:
        person["chart_context"].append(context)
    ref = f"{source}:{source_ref}"
    if ref not in person["source_refs"]:
        person["source_refs"].append(ref)
    if review_status != "confirmed":
        person["review_status"] = "needs_review"


def _header_index(rows: list[list[str]]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:20]):
        normalized = [value.casefold() for value in row]
        name = next((i for i, value in enumerate(normalized)
                     if value in {"name", "name of employee or applicant", "employee name"}), None)
        role = next((i for i, value in enumerate(normalized) if value in {"role", "position"}), None)
        if name is not None and role is not None:
            personnel = next((i for i, value in enumerate(normalized)
                              if "personnel" in value and ("number" in value or "id" in value)), None)
            discipline = next((i for i, value in enumerate(normalized)
                               if value in {"discipline", "disciplines", "category"}), None)
            return index, {"name": name, "role": role, "employee_id": personnel,
                           "context": discipline}
    return None


def _excel(path: Path, store: dict) -> None:
    year, asset = _year(path), _asset(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [[_clean(value) for value in row] for row in ws.iter_rows(values_only=True)]
        explicit = _header_index(rows)
        if explicit:
            header_row, columns = explicit
            for offset, row in enumerate(rows[header_row + 1:], header_row + 2):
                if columns["name"] >= len(row):
                    continue
                name = row[columns["name"]]
                title = row[columns["role"]] if columns["role"] < len(row) else ""
                employee = (row[columns["employee_id"]] if columns["employee_id"] is not None
                            and columns["employee_id"] < len(row) else "")
                context = (row[columns["context"]] if columns["context"] is not None
                           and columns["context"] < len(row) else ws.title)
                _add(store, year=year, asset=asset, name=name, title=title, context=context,
                     employee_id=employee, source=path.name, source_ref=f"{ws.title}!row{offset}")
            continue

        # Matrix charts: an IPT/workstream heading governs the following names in that column until a
        # new heading.  The heading is context, never a personal title.
        contexts: dict[int, str] = {}
        for row_index, row in enumerate(rows, 1):
            for column, value in enumerate(row, 1):
                if not value:
                    continue
                if _HEADING_WORDS.search(value) and not _looks_like_name(value):
                    contexts[column] = value
                    continue
                context = contexts.get(column, "")
                if not context:
                    continue
                _add(store, year=year, asset=asset, name=value, context=context,
                     source=path.name, source_ref=f"{ws.title}!{ws.cell(row_index, column).coordinate}",
                     review_status="confirmed" if _looks_like_name(value) else "needs_review")
    wb.close()


def _pptx(path: Path, store: dict) -> None:
    year, asset = _year(path), _asset(path)
    prs = Presentation(path)
    for slide_number, slide in enumerate(prs.slides, 1):
        texts = [str(getattr(shape, "text", "")).strip() for shape in slide.shapes
                 if str(getattr(shape, "text", "")).strip()]
        context = next((_clean(text) for text in texts
                        if len(_clean(text)) < 100 and not _looks_like_name(_clean(text))),
                       f"Slide {slide_number}")
        for shape_index, text in enumerate(texts, 1):
            lines = [_clean(line) for line in re.split(r"[\n/]", text) if _clean(line)]
            if not lines:
                continue
            name = lines[0]
            if not _looks_like_name(name):
                continue
            title = " / ".join(lines[1:])
            _add(store, year=year, asset=asset, name=name, title=title, context=context,
                 source=path.name, source_ref=f"slide{slide_number}:shape{shape_index}",
                 review_status="confirmed" if title else "needs_review")


def build() -> dict:
    years: dict[str, Any] = {}
    for path in sorted(SOURCE_ROOT.rglob("*")):
        if path.suffix.casefold() == ".xlsx":
            _excel(path, years)
        elif path.suffix.casefold() == ".pptx":
            _pptx(path, years)
    for year_data in years.values():
        for asset_data in year_data.values():
            files = sorted({ref.split(":", 1)[0] for person in asset_data["people"].values()
                            for ref in person["source_refs"]})
            asset_data["sources"] = [{"file": file, "effective_date": ""} for file in files]
            asset_data["people"] = sorted(asset_data["people"].values(), key=lambda row: row["name"].casefold())
    return {"schema_version": 1, "audit_years": years}


if __name__ == "__main__":
    payload = build()
    OUTPUT.write_text(yaml.safe_dump(payload, sort_keys=False, allow_unicode=True), encoding="utf-8")
    print(f"wrote {OUTPUT}")
