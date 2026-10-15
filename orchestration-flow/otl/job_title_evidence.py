"""Deterministic staff/catalog/IPT evidence and batched structured job-title analysis."""
from __future__ import annotations

import contextvars
import json
import re
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from functools import lru_cache
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from rapidfuzz.fuzz import token_set_ratio

from otl import llm_adapter as llm
from otl.ipt_registry import status as ipt_status
from otl.jade_rules import exempt_terms


TARGET_TOKENS = 180_000
MAX_TOKENS = 220_000
MAX_JOBS = 4
MAX_CANDIDATES = 8
CATALOG_SHEET = "Job Catalogue"


class StaffTitlePeriod(BaseModel):
    employee_id: str = ""
    adopted_name: str
    position: str
    explicit_job_profile: str | None = None
    cost_center_code: str | None = None
    cost_center_name: str | None = None
    manager_id: str | None = None
    manager_name: str | None = None
    months: list[str] = Field(default_factory=list)
    first_month: str = ""
    last_month: str = ""
    source_rows: list[int] = Field(default_factory=list)


class CatalogCandidate(BaseModel):
    job_family_group: str = ""
    job_family: str = ""
    job_profile: str
    job_code: str = ""
    management_level: str = ""
    job_summary: str = ""
    job_description: str = ""
    job_profile_exempt: str = ""
    match_score: float = 0.0
    match_basis: str = "fuzzy_position"


class JobTitlePersonPacket(BaseModel):
    evidence_id: str
    person: str
    identity_status: str
    staff_title_periods: list[StaffTitlePeriod] = Field(default_factory=list)
    catalog_candidates_by_period: dict[str, list[CatalogCandidate]] = Field(default_factory=dict)
    ipt_registry_match: dict[str, Any] | None = None
    staff_source: str = ""
    catalog_source: str = ""


class JobTitlePersonFinding(BaseModel):
    person: str
    evidence_ids_consumed: list[str] = Field(default_factory=list)
    adopted_positions: list[str] = Field(default_factory=list)
    selected_job_profiles: list[str] = Field(default_factory=list)
    selected_job_profiles_by_period: dict[str, list[str]] = Field(default_factory=dict)
    technical_status: Literal["technical", "non_technical", "unclear"] = "unclear"
    technical_status_by_period: dict[str, Literal["technical", "non_technical", "unclear"]] = Field(
        default_factory=dict)
    technical_rationale: str = ""
    job_profile_status: Literal["matched", "no_match", "multiple_close", "unclear"] = "unclear"
    needs_raw_review: bool = False
    raw_review_reason: str = ""


class JobTitleDecision(BaseModel):
    findings: list[JobTitlePersonFinding] = Field(default_factory=list)


class PreparedJobTitleEvidence(BaseModel):
    packets: list[JobTitlePersonPacket] = Field(default_factory=list)
    unresolved_reasons: dict[str, list[str]] = Field(default_factory=dict)
    staff_source: str = ""
    catalog_source: str = ""


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _clean(value).casefold()).strip()


def _clip(value: str, limit: int = 3500) -> str:
    text = _clean(value)
    if len(text) <= limit:
        return text
    cut = text.rfind(". ", 0, limit)
    return text[:cut + 1 if cut > limit // 2 else limit].rstrip()


def _month_key(value: str) -> tuple[int, str]:
    months = {name.casefold(): i for i, name in enumerate(
        ("January", "February", "March", "April", "May", "June", "July", "August",
         "September", "October", "November", "December"), 1)}
    text = _clean(value)
    return months.get(text.casefold(), 99), text


def _first_role_value(row: dict, entry: dict, role: str, *, exclude_id: bool = False) -> str:
    columns = (entry.get("columns") or {}).get(role) or {}
    for label in columns:
        if exclude_id and "id" in label.casefold():
            continue
        value = _clean(row.get(label))
        if value:
            return value
    return ""


def _jade_employee_ids(state: dict) -> dict[str, str]:
    jade = state.get("jade_analyzer") or {}
    if hasattr(jade, "model_dump"):
        jade = jade.model_dump()
    out: dict[str, str] = {}
    for row in jade.get("ledger_rows") or []:
        if hasattr(row, "model_dump"):
            row = row.model_dump()
        person, employee = _clean(row.get("person")), re.sub(r"\D", "", str(row.get("employee_id") or ""))
        if person and employee:
            out.setdefault(person, employee)
    return out


def title_periods(index: dict, people: list[str], state: dict) -> tuple[
    dict[str, list[StaffTitlePeriod]], dict[str, list[str]], str
]:
    periods: dict[str, list[StaffTitlePeriod]] = {}
    reasons: dict[str, list[str]] = defaultdict(list)
    source_name = ""
    jade_ids = _jade_employee_ids(state)
    staff_sources = [source for source in index.get("sources") or [] if source.get("role") == "staff"]
    for source in staff_sources:
        entries = source.get("entries") or []
        source_name = source_name or ", ".join(dict.fromkeys(e.get("name", "") for e in entries))
        selected = source.get("selected") or {}
        for person in people:
            matches = selected.get(person) or []
            rows: list[tuple[dict, dict, str]] = []
            for match in matches:
                name = match.get("name_in_file", "")
                for entry in entries:
                    for row in (entry.get("row_data") or {}).get(name) or []:
                        rows.append((row, entry, name))
            wanted_id = jade_ids.get(person, "")
            if wanted_id:
                matching_id = [item for item in rows if re.sub(
                    r"\D", "", _first_role_value(item[0], item[1], "person_key")) == wanted_id]
                if matching_id:
                    rows = matching_id
            if not rows:
                reasons[person].append("no selected staff-listing row")
                periods[person] = []
                continue
            groups: dict[tuple, dict[str, Any]] = {}
            for row, entry, adopted in rows:
                employee = _first_role_value(row, entry, "person_key")
                position = _first_role_value(row, entry, "position", exclude_id=True)
                profile = _first_role_value(row, entry, "job_profile")
                manager_id = _first_role_value(row, entry, "reports_to")
                manager_name = _first_role_value(row, entry, "manager_name")
                month = _first_role_value(row, entry, "period")
                cost_code = next((_clean(row.get(key)) for key in row
                                  if "cost center" in key.casefold() and ("id" in key.casefold()
                                  or "key" in key.casefold()) and _clean(row.get(key))), "")
                cost_name = next((_clean(row.get(key)) for key in row
                                  if "cost center" in key.casefold() and not ("id" in key.casefold()
                                  or "key" in key.casefold()) and _clean(row.get(key))), "")
                key = (employee, adopted, position, profile, cost_code, cost_name, manager_id, manager_name)
                item = groups.setdefault(key, {"months": [], "rows": []})
                if month and month not in item["months"]:
                    item["months"].append(month)
                source_row = int(row.get("_row") or 0) + 1
                if source_row and source_row not in item["rows"]:
                    item["rows"].append(source_row)
            built: list[StaffTitlePeriod] = []
            for key, values in groups.items():
                employee, adopted, position, profile, cost_code, cost_name, manager_id, manager_name = key
                months = sorted(values["months"], key=_month_key)
                built.append(StaffTitlePeriod(
                    employee_id=employee, adopted_name=adopted, position=position,
                    explicit_job_profile=profile or None, cost_center_code=cost_code or None,
                    cost_center_name=cost_name or None, manager_id=manager_id or None,
                    manager_name=manager_name or None, months=months,
                    first_month=months[0] if months else "", last_month=months[-1] if months else "",
                    source_rows=sorted(values["rows"]),
                ))
            periods[person] = sorted(built, key=lambda item: (_month_key(item.first_month), item.position))
            identities = {period.adopted_name for period in built}
            employee_ids = {period.employee_id for period in built if period.employee_id}
            if len(identities) > 1 and len(employee_ids) > 1:
                reasons[person].append("multiple staff identities with different employee IDs")
    return periods, dict(reasons), source_name


def _catalog_key(path: str) -> tuple[str, int, int]:
    stat = Path(path).stat()
    return str(Path(path).resolve()), stat.st_size, stat.st_mtime_ns


@lru_cache(maxsize=4)
def _catalog_rows(path: str, _size: int, _mtime: int) -> tuple[CatalogCandidate, ...]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if CATALOG_SHEET not in wb.sheetnames:
            raise ValueError(f"worksheet {CATALOG_SHEET!r} not found")
        ws = wb[CATALOG_SHEET]
        rows = ws.iter_rows(values_only=True)
        header = [_clean(value) for value in next(rows)]
        required = ["Job Family Group", "Job Family", "Job Profile", "Job Code", "Management Level",
                    "Job Summary", "Job Description", "Job Profile Exempt"]
        missing = [name for name in required if name not in header]
        if missing:
            raise ValueError("catalogue columns missing: " + ", ".join(missing))
        indices = {name: header.index(name) for name in required}
        out: list[CatalogCandidate] = []
        for row in rows:
            profile = _clean(row[indices["Job Profile"]] if indices["Job Profile"] < len(row) else "")
            if not profile:
                continue
            out.append(CatalogCandidate(
                job_family_group=_clean(row[indices["Job Family Group"]]),
                job_family=_clean(row[indices["Job Family"]]), job_profile=profile,
                job_code=_clean(row[indices["Job Code"]]),
                management_level=_clean(row[indices["Management Level"]]),
                job_summary=_clip(row[indices["Job Summary"]]),
                job_description=_clip(row[indices["Job Description"]]),
                job_profile_exempt=_clean(row[indices["Job Profile Exempt"]]),
            ))
        return tuple(out)
    finally:
        wb.close()


def catalog_rows(path: str) -> tuple[CatalogCandidate, ...]:
    resolved, size, mtime = _catalog_key(path)
    return _catalog_rows(resolved, size, mtime)


def catalog_candidates(period: StaffTitlePeriod, rows: tuple[CatalogCandidate, ...]) -> list[CatalogCandidate]:
    explicit = _norm(period.explicit_job_profile or "")
    if explicit:
        exact = [row.model_copy(update={"match_score": 100.0, "match_basis": "explicit_job_profile"})
                 for row in rows if _norm(row.job_profile) == explicit]
        if exact:
            return exact[:MAX_CANDIDATES]
    position = _norm(period.position)
    exact = [row.model_copy(update={"match_score": 100.0, "match_basis": "exact_position"})
             for row in rows if _norm(row.job_profile) == position]
    if exact:
        return exact[:MAX_CANDIDATES]
    scored = sorted(
        ((token_set_ratio(position, _norm(row.job_profile)), row) for row in rows),
        key=lambda pair: pair[0], reverse=True,
    )
    if not scored:
        return []
    cutoff_index = min(4, len(scored) - 1)
    cutoff = scored[cutoff_index][0]
    selected = [(score, row) for score, row in scored if score >= cutoff][:MAX_CANDIDATES]
    return [row.model_copy(update={"match_score": round(score, 2), "match_basis": "fuzzy_position"})
            for score, row in selected]


def prepare_evidence(*, state: dict, index: dict, people: list[str], staff_listing: str,
                     jobs_catalogue: str) -> PreparedJobTitleEvidence:
    periods, unresolved, staff_source = title_periods(index, people, state)
    catalogue = catalog_rows(jobs_catalogue)
    case = state.get("case") or {}
    audit_year, asset = str(case.get("year") or ""), str(case.get("asset") or "")
    packets: list[JobTitlePersonPacket] = []
    for position, person in enumerate(people, 1):
        person_periods = periods.get(person) or []
        candidates = {
            f"period-{i:02d}": catalog_candidates(period, catalogue)
            for i, period in enumerate(person_periods, 1)
        }
        employee_id = next((period.employee_id for period in person_periods if period.employee_id), "")
        ipt_value, ipt_match = ipt_status(
            audit_year=audit_year, asset=asset, person=person, employee_id=employee_id)
        if ipt_match:
            ipt_match = dict(ipt_match) | {"status": ipt_value, "asset": asset, "audit_year": audit_year}
        packets.append(JobTitlePersonPacket(
            evidence_id=f"job-title:{position:04d}", person=person,
            identity_status="resolved" if person_periods and not unresolved.get(person) else "unresolved",
            staff_title_periods=person_periods, catalog_candidates_by_period=candidates,
            ipt_registry_match=ipt_match, staff_source=staff_source or Path(staff_listing).name,
            catalog_source=f"{Path(jobs_catalogue).name}!{CATALOG_SHEET}",
        ))
    return PreparedJobTitleEvidence(
        packets=packets, unresolved_reasons=unresolved,
        staff_source=staff_source or Path(staff_listing).name,
        catalog_source=f"{Path(jobs_catalogue).name}!{CATALOG_SHEET}",
    )


def _packet_tokens(packet: JobTitlePersonPacket) -> int:
    return llm.estimate_text_tokens(packet.model_dump_json(exclude_none=True))


def build_shards(packets: list[JobTitlePersonPacket]) -> tuple[list[list[JobTitlePersonPacket]], dict[str, list[str]]]:
    shards: list[list[JobTitlePersonPacket]] = []
    reasons: dict[str, list[str]] = defaultdict(list)
    current: list[JobTitlePersonPacket] = []
    tokens = 0
    for packet in packets:
        size = _packet_tokens(packet)
        if size > MAX_TOKENS:
            reasons[packet.person].append("normalized person packet exceeds the hard token limit")
            continue
        if current and tokens + size > TARGET_TOKENS:
            shards.append(current); current, tokens = [], 0
        current.append(packet); tokens += size
    if current:
        shards.append(current)
    return shards, dict(reasons)


_SYSTEM = """You classify each supplied person's documented job role.

Use only the supplied staff title periods and catalogue candidates. Select profiles only from the
candidate lists. Judge technical in nature from documented duties: engineering, geological,
geophysical, technical, operational, planning, scheduling, maintenance, turnaround, drilling,
completion, construction, or other special/specific operational skills are technical; dedicated
administrative, secretarial, or generic business support is non-technical. Analyze every title period.
Populate technical_status_by_period and selected_job_profiles_by_period using every supplied period ID.
If periods materially conflict, return unclear and request raw review. Never change IPT membership,
employee identity, source titles, or source rows. Return every evidence ID exactly once."""


def _analyze_one(packets: list[JobTitlePersonPacket], context: str) -> JobTitleDecision:
    user = (
        f"CASE CONTEXT:\n{context}\n\nNORMALIZED JOB-TITLE EVIDENCE:\n"
        + json.dumps([packet.model_dump(mode="json", exclude_none=True) for packet in packets],
                     ensure_ascii=False, separators=(",", ":"))
    )
    return llm.structured("job_title", _SYSTEM, user, JobTitleDecision)


def analyze_packets(packets: list[JobTitlePersonPacket], *, context: str = "") -> tuple[
    list[JobTitlePersonFinding], dict[str, list[str]], int
]:
    shards, reasons = build_shards(packets)
    if not shards:
        return [], reasons, 0
    jobs = min(MAX_JOBS, len(shards))
    if jobs == 1:
        decisions = [_analyze_one(shards[0], context)]
    else:
        with ThreadPoolExecutor(max_workers=jobs, thread_name_prefix="job-title") as pool:
            futures = [pool.submit(contextvars.copy_context().run, _analyze_one, shard, context)
                       for shard in shards]
            decisions = [future.result() for future in futures]
    packet_by_person = {packet.person: packet for packet in packets}
    findings: list[JobTitlePersonFinding] = []
    seen_evidence: set[str] = set()
    for decision in decisions:
        for finding in decision.findings:
            packet = packet_by_person.get(finding.person)
            if not packet:
                continue
            expected = packet.evidence_id
            if finding.evidence_ids_consumed != [expected] or expected in seen_evidence:
                reasons.setdefault(finding.person, []).append("structured evidence coverage failed")
                continue
            seen_evidence.add(expected)
            allowed_profiles = {
                candidate.job_profile for values in packet.catalog_candidates_by_period.values()
                for candidate in values
            }
            if any(profile not in allowed_profiles for profile in finding.selected_job_profiles):
                reasons.setdefault(finding.person, []).append("structured result selected an unknown profile")
                continue
            period_keys = set(packet.catalog_candidates_by_period)
            if len(period_keys) > 1:
                if set(finding.technical_status_by_period) != period_keys:
                    reasons.setdefault(finding.person, []).append(
                        "structured result omitted a title-period conclusion")
                    continue
                period_profiles_valid = all(
                    profile in {candidate.job_profile for candidate in
                                packet.catalog_candidates_by_period.get(period, [])}
                    for period, profiles in finding.selected_job_profiles_by_period.items()
                    for profile in profiles
                )
                if not period_profiles_valid:
                    reasons.setdefault(finding.person, []).append(
                        "structured result selected an unknown period profile")
                    continue
                period_statuses = set(finding.technical_status_by_period.values()) - {"unclear"}
                if len(period_statuses) > 1:
                    finding = finding.model_copy(update={
                        "technical_status": "unclear", "needs_raw_review": True,
                        "raw_review_reason": "technical status conflicts across title periods",
                    })
            if finding.needs_raw_review:
                reasons.setdefault(finding.person, []).append(
                    finding.raw_review_reason or "structured analysis requested raw review")
            findings.append(finding)
    for packet in packets:
        if packet.evidence_id not in seen_evidence:
            reasons.setdefault(packet.person, []).append("structured result omitted this person")
    return findings, reasons, len(shards)


def exempt_status(*, positions: list[str], profiles: list[str], ipt_match: dict | None) -> str:
    values = [*positions, *profiles]
    if ipt_match and ipt_match.get("registry_title"):
        values.append(str(ipt_match["registry_title"]))
    haystack = set(_norm(" ".join(values)).split())
    return "exempt" if any(set(_norm(term).split()) <= haystack for term in exempt_terms()) else "non_exempt"


__all__ = [
    "CatalogCandidate", "JobTitleDecision", "JobTitlePersonFinding", "JobTitlePersonPacket",
    "PreparedJobTitleEvidence", "StaffTitlePeriod", "analyze_packets", "build_shards",
    "catalog_candidates", "catalog_rows", "exempt_status", "prepare_evidence", "title_periods",
]
