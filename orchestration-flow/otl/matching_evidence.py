"""Deterministic CFR workbook normalization and batched per-person mapping selection."""
from __future__ import annotations

import contextvars
import hashlib
import json
import re
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from functools import lru_cache
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from rapidfuzz.fuzz import token_set_ratio

from otl import llm_adapter as llm


TARGET_TOKENS = 180_000
MAX_TOKENS = 220_000
MAX_JOBS = 4
MAX_CANDIDATES = 8


class CfrMappingCandidate(BaseModel):
    candidate_id: str
    job_profile_description: str
    job_title: str = ""
    job_summary: str = ""
    job_description: str = ""
    cfr_primary: str = ""
    cfr_secondary: str = ""
    cfr_additional: str = ""
    subpart_labels: str = ""
    rationale: str = ""
    source_reference: str = ""
    notes: str = ""
    row_count: int = 1
    match_score: float = 0.0
    match_basis: str = ""


class MatchingPersonPacket(BaseModel):
    evidence_id: str
    person: str
    positions: list[str] = Field(default_factory=list)
    selected_job_profiles: list[str] = Field(default_factory=list)
    job_summaries: dict[str, str] = Field(default_factory=dict)
    job_descriptions: dict[str, str] = Field(default_factory=dict)
    candidates: list[CfrMappingCandidate] = Field(default_factory=list)


class MatchingPersonFinding(BaseModel):
    person: str
    evidence_ids_consumed: list[str] = Field(default_factory=list)
    selected_candidate_ids: list[str] = Field(default_factory=list)
    mapping_status: Literal["matched", "no_match", "multiple_close", "unclear"] = "unclear"
    rationale: str = ""
    needs_raw_review: bool = False
    raw_review_reason: str = ""


class MatchingDecision(BaseModel):
    findings: list[MatchingPersonFinding] = Field(default_factory=list)


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _norm(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _clean(value).casefold()).strip()


def _clip(value: str, limit: int = 2500) -> str:
    text = _clean(value)
    if len(text) <= limit:
        return text
    cut = text.rfind(". ", 0, limit)
    return text[:cut + 1 if cut > limit // 2 else limit].rstrip()


def _cache_key(path: str) -> tuple[str, int, int]:
    resolved = Path(path).resolve(); stat = resolved.stat()
    return str(resolved), stat.st_size, stat.st_mtime_ns


@lru_cache(maxsize=4)
def _mapping_rows(path: str, _size: int, _mtime: int) -> tuple[CfrMappingCandidate, ...]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [_clean(value) for value in next(rows)]
        required = ["Job Title", "Job Profile Description", "Job summary 1", "Job description 1",
                    "CFR Primary", "CFR Secondary", "CFR Additional", "Subpart Labels", "Rationale",
                    "Source Reference", "Notes"]
        missing = [field for field in required if field not in header]
        if missing:
            raise ValueError("CFR mapping columns missing: " + ", ".join(missing))
        indices = {field: header.index(field) for field in required}
        grouped: dict[tuple, dict] = {}
        for row_number, row in enumerate(rows, 2):
            values = tuple(_clean(row[indices[field]] if indices[field] < len(row) else "")
                           for field in required)
            if not values[1]:
                continue
            item = grouped.setdefault(values, {"count": 0, "rows": []})
            item["count"] += 1
            if len(item["rows"]) < 5:
                item["rows"].append(row_number)
        out = []
        for values, meta in sorted(grouped.items(), key=lambda pair: (_norm(pair[0][1]), _norm(pair[0][0]))):
            digest = hashlib.sha1("|".join(values).encode("utf-8")).hexdigest()[:12]
            out.append(CfrMappingCandidate(
                candidate_id=f"cfr-{digest}", job_title=values[0],
                job_profile_description=values[1], job_summary=_clip(values[2]),
                job_description=_clip(values[3]), cfr_primary=values[4], cfr_secondary=values[5],
                cfr_additional=values[6], subpart_labels=values[7], rationale=_clip(values[8]),
                source_reference=values[9], notes=_clip(values[10], 1000), row_count=meta["count"],
            ))
        return tuple(out)
    finally:
        wb.close()


def mapping_rows(path: str) -> tuple[CfrMappingCandidate, ...]:
    resolved, size, mtime = _cache_key(path)
    return _mapping_rows(resolved, size, mtime)


def candidates_for_person(*, positions: list[str], profiles: list[str], rows: tuple[CfrMappingCandidate, ...]) -> list[CfrMappingCandidate]:
    profile_norms = {_norm(value) for value in profiles if _norm(value)}
    exact = [row for row in rows if _norm(row.job_profile_description) in profile_norms]
    pool = exact or list(rows)
    queries = [*profiles, *positions]
    scored = []
    for row in pool:
        score = max((token_set_ratio(_norm(query), _norm(row.job_profile_description + " " + row.job_title))
                     for query in queries if _norm(query)), default=0.0)
        basis = "exact_job_profile" if row in exact else "fuzzy_role_profile"
        scored.append((score, row.model_copy(update={"match_score": round(score, 2), "match_basis": basis})))
    scored.sort(key=lambda pair: pair[0], reverse=True)
    if not scored:
        return []
    cutoff = scored[min(4, len(scored) - 1)][0]
    return [row for score, row in scored if score >= cutoff][:MAX_CANDIDATES]


def prepare_packets(state: dict, mapping_path: str) -> list[MatchingPersonPacket]:
    job = state.get("job_title") or {}
    if hasattr(job, "model_dump"):
        job = job.model_dump()
    people = list((job.get("positions_by_person") or {}).keys()
                  or (job.get("identified_names_staff_listing") or {}).keys())
    rows = mapping_rows(mapping_path)
    packets = []
    for index, person in enumerate(people, 1):
        positions = list((job.get("positions_by_person") or {}).get(person) or [])
        if not positions and (job.get("identified_names_staff_listing") or {}).get(person):
            positions = [str(job["identified_names_staff_listing"][person])]
        profiles = list((job.get("selected_job_profiles_by_person") or {}).get(person) or [])
        summaries = {profile: str((job.get("identified_job_data") or {}).get(profile) or "")
                     for profile in profiles}
        descriptions = {profile: str((job.get("identified_job_descriptions") or {}).get(profile) or "")
                        for profile in profiles}
        packets.append(MatchingPersonPacket(
            evidence_id=f"matching:{index:04d}", person=person, positions=positions,
            selected_job_profiles=profiles, job_summaries=summaries, job_descriptions=descriptions,
            candidates=candidates_for_person(positions=positions, profiles=profiles, rows=rows),
        ))
    return packets


def _packet_tokens(packet: MatchingPersonPacket) -> int:
    return llm.estimate_text_tokens(packet.model_dump_json(exclude_none=True))


def build_shards(packets: list[MatchingPersonPacket]) -> tuple[list[list[MatchingPersonPacket]], dict[str, list[str]]]:
    shards: list[list[MatchingPersonPacket]] = []
    reasons: dict[str, list[str]] = defaultdict(list)
    current: list[MatchingPersonPacket] = []; tokens = 0
    for packet in packets:
        size = _packet_tokens(packet)
        if size > MAX_TOKENS:
            reasons[packet.person].append("normalized CFR packet exceeds the hard token limit")
            continue
        if current and tokens + size > TARGET_TOKENS:
            shards.append(current); current, tokens = [], 0
        current.append(packet); tokens += size
    if current:
        shards.append(current)
    return shards, dict(reasons)


_SYSTEM = """You map documented job roles onto the supplied CFR workbook candidates.

Use the person's staff positions, selected job profiles, job summaries and job descriptions to select
only the candidate IDs that represent the same role/profile. Preserve all supplied CFR text exactly.
Do not decide whether labor is chargeable, billable, technical, onsite, offsite, or covered by overhead.
Do not apply settlement, accounting-procedure, or OTL rules. Return every evidence ID exactly once. A
single clearly supported candidate is matched; indistinguishable candidates are multiple_close; no
supported candidate is no_match; missing role evidence is unclear."""


def _one(shard: list[MatchingPersonPacket], context: str) -> MatchingDecision:
    user = f"CASE CONTEXT:\n{context}\n\nNORMALIZED CFR CANDIDATES:\n" + json.dumps(
        [packet.model_dump(mode="json") for packet in shard], ensure_ascii=False, separators=(",", ":"))
    return llm.structured("matching", _SYSTEM, user, MatchingDecision)


def analyze_packets(packets: list[MatchingPersonPacket], *, context: str = "") -> tuple[
    list[MatchingPersonFinding], dict[str, list[str]], int
]:
    deterministic: list[MatchingPersonFinding] = []
    ambiguous: list[MatchingPersonPacket] = []
    for packet in packets:
        if len(packet.candidates) == 1:
            deterministic.append(MatchingPersonFinding(
                person=packet.person, evidence_ids_consumed=[packet.evidence_id],
                selected_candidate_ids=[packet.candidates[0].candidate_id], mapping_status="matched",
                rationale="One deterministic CFR mapping candidate remained after profile/title normalization.",
            ))
        elif not packet.candidates:
            deterministic.append(MatchingPersonFinding(
                person=packet.person, evidence_ids_consumed=[packet.evidence_id], mapping_status="no_match",
                rationale="No CFR mapping candidate was supported by the resolved role/profile.",
            ))
        else:
            ambiguous.append(packet)
    shards, reasons = build_shards(ambiguous)
    if shards:
        jobs = min(MAX_JOBS, len(shards))
        if jobs == 1:
            decisions = [_one(shards[0], context)]
        else:
            with ThreadPoolExecutor(max_workers=jobs, thread_name_prefix="matching") as pool:
                futures = [pool.submit(contextvars.copy_context().run, _one, shard, context)
                           for shard in shards]
                decisions = [future.result() for future in futures]
        packet_map = {packet.person: packet for packet in ambiguous}
        seen: set[str] = set()
        for decision in decisions:
            for finding in decision.findings:
                packet = packet_map.get(finding.person)
                if (not packet or finding.evidence_ids_consumed != [packet.evidence_id]
                        or packet.evidence_id in seen):
                    reasons.setdefault(finding.person, []).append("structured CFR coverage failed")
                    continue
                seen.add(packet.evidence_id)
                allowed = {candidate.candidate_id for candidate in packet.candidates}
                if any(candidate not in allowed for candidate in finding.selected_candidate_ids):
                    reasons.setdefault(finding.person, []).append("structured result selected unknown CFR candidate")
                    continue
                if finding.needs_raw_review:
                    reasons.setdefault(finding.person, []).append(
                        finding.raw_review_reason or "structured CFR analysis requested raw review")
                deterministic.append(finding)
        for packet in ambiguous:
            if packet.evidence_id not in seen:
                reasons.setdefault(packet.person, []).append("structured CFR result omitted this person")
    return deterministic, reasons, len(shards)


__all__ = ["CfrMappingCandidate", "MatchingDecision", "MatchingPersonFinding", "MatchingPersonPacket",
           "analyze_packets", "build_shards", "candidates_for_person", "mapping_rows", "prepare_packets"]
